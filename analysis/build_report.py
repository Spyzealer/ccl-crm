import json, datetime
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
from openpyxl.chart import BarChart, PieChart, Reference

d = json.load(open('/tmp/gs_data.json'))
customers = d['customers']
services = d['services']
followups = d['followups']
churned = d['churned']

wb = Workbook()

HEADER_FILL = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
HEADER_FONT = Font(bold=True, color="FFFFFF")

def write_table(ws, headers, rows):
    ws.append(headers)
    for c in range(1, len(headers)+1):
        cell = ws.cell(row=1, column=c)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(horizontal="center")
    for r in rows:
        ws.append(r)
    ws.freeze_panes = "A2"
    for c in range(1, len(headers)+1):
        maxlen = max([len(str(headers[c-1]))] + [len(str(r[c-1])) if c-1 < len(r) else 0 for r in rows])
        ws.column_dimensions[get_column_letter(c)].width = min(max(maxlen+2, 10), 45)
    if rows:
        ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{len(rows)+1}"

# ---------- Sheet: Customers ----------
ws = wb.active
ws.title = "Customers"
headers = ['customer','group_name','tier','tier_remark','contact','tel','email','note','updated_at','updated_by']
rows = [[c.get(h,'') for h in headers] for c in customers]
write_table(ws, headers, rows)

# ---------- Sheet: Services ----------
ws = wb.create_sheet("Services")
headers = ['id','customer','item','type','qty','price_unit','price_year','start','expire','contact','tel','email','note','source','updated_at','updated_by']
rows = [[s.get(h,'') for h in headers] for s in services]
write_table(ws, headers, rows)

# ---------- Sheet: Followups ----------
ws = wb.create_sheet("Followups")
headers = ['id','customer','date','type','note','owner','status','updated_at','updated_by']
rows = [[f.get(h,'') for h in headers] for f in followups]
write_table(ws, headers, rows)

# ---------- Sheet: Churned ----------
ws = wb.create_sheet("Churned")
headers = ['name','services','lost_revenue','expire','reason','winback','updated_at']
rows = [[c.get(h,'') for h in headers] for c in churned]
write_table(ws, headers, rows)

# ---------- Analysis: Summary ----------
ws = wb.create_sheet("Summary", 0)
ws.append(["CCL CRM — สรุปข้อมูลเพื่อวิเคราะห์", ""])
ws['A1'].font = Font(bold=True, size=14)
ws.append([])

total_customers = len(customers)
total_services = len(services)
total_revenue = sum([s.get('price_year') or 0 for s in services if isinstance(s.get('price_year'), (int,float))])
total_lost_revenue = sum([c.get('lost_revenue') or 0 for c in churned if isinstance(c.get('lost_revenue'), (int,float))])
open_followups = len([f for f in followups if str(f.get('status','')).lower()=='open'])
tier_counts = {}
for c in customers:
    t = c.get('tier') or 'ไม่ระบุ'
    tier_counts[t] = tier_counts.get(t,0)+1

type_counts = {}
type_revenue = {}
for s in services:
    t = s.get('type') or 'ไม่ระบุ'
    type_counts[t] = type_counts.get(t,0)+1
    type_revenue[t] = type_revenue.get(t,0) + (s.get('price_year') or 0 if isinstance(s.get('price_year'),(int,float)) else 0)

# Expiring soon (next 90 days)
now = datetime.datetime.utcnow()
def parse_dt(s):
    if not s: return None
    try:
        return datetime.datetime.fromisoformat(s.replace('Z','+00:00')).replace(tzinfo=None)
    except Exception:
        return None
expiring_soon = []
for s in services:
    dt = parse_dt(s.get('expire'))
    if dt:
        days = (dt - now).days
        if 0 <= days <= 90:
            expiring_soon.append((s, days))
expiring_soon.sort(key=lambda x: x[1])

kpi_rows = [
    ["จำนวนลูกค้าทั้งหมด", total_customers],
    ["จำนวนบริการทั้งหมด", total_services],
    ["รายได้รวมต่อปี (จากบริการที่ยังใช้อยู่)", total_revenue],
    ["รายได้ที่เสียไป (Churned)", total_lost_revenue],
    ["Follow-up ที่ยังเปิดอยู่ (Open)", open_followups],
    ["จำนวนลูกค้าที่เลิกใช้ (Churned)", len(churned)],
    ["บริการที่จะหมดอายุใน 90 วันข้างหน้า", len(expiring_soon)],
]
ws.append(["KPI", "ค่า"])
for i in range(3,4): pass
r0 = ws.max_row+1
for row in kpi_rows:
    ws.append(row)
for c in range(1,3):
    ws.cell(row=r0-1, column=c).font = HEADER_FONT
    ws.cell(row=r0-1, column=c).fill = HEADER_FILL

ws.append([])
ws.append(["สัดส่วนลูกค้าตาม Tier"])
ws.cell(row=ws.max_row, column=1).font = Font(bold=True)
tier_start = ws.max_row+1
ws.append(["Tier","จำนวนลูกค้า"])
for c in range(1,3):
    ws.cell(row=ws.max_row, column=c).font = HEADER_FONT
    ws.cell(row=ws.max_row, column=c).fill = HEADER_FILL
for t,cnt in sorted(tier_counts.items(), key=lambda x:-x[1]):
    ws.append([t if t else "ไม่ระบุ", cnt])
tier_end = ws.max_row

ws.append([])
ws.append(["รายได้ต่อปีแยกตามประเภทบริการ"])
ws.cell(row=ws.max_row, column=1).font = Font(bold=True)
type_start = ws.max_row+1
ws.append(["ประเภทบริการ","จำนวนรายการ","รายได้ต่อปี (บาท)"])
for c in range(1,4):
    ws.cell(row=ws.max_row, column=c).font = HEADER_FONT
    ws.cell(row=ws.max_row, column=c).fill = HEADER_FILL
for t,cnt in sorted(type_counts.items(), key=lambda x:-type_revenue.get(x[0],0)):
    ws.append([t if t else "ไม่ระบุ", cnt, type_revenue.get(t,0)])
type_end = ws.max_row

for c in range(1,4):
    ws.column_dimensions[get_column_letter(c)].width = 32

# Charts
pie = PieChart()
pie.title = "สัดส่วนลูกค้าตาม Tier"
data = Reference(ws, min_col=2, min_row=tier_start, max_row=tier_end)
cats = Reference(ws, min_col=1, min_row=tier_start+1, max_row=tier_end)
pie.add_data(data, titles_from_data=True)
pie.set_categories(cats)
ws.add_chart(pie, "E3")

bar = BarChart()
bar.title = "รายได้ต่อปีแยกตามประเภทบริการ"
bar.y_axis.title = "บาท/ปี"
data2 = Reference(ws, min_col=3, min_row=type_start, max_row=type_end)
cats2 = Reference(ws, min_col=1, min_row=type_start+1, max_row=type_end)
bar.add_data(data2, titles_from_data=True)
bar.set_categories(cats2)
ws.add_chart(bar, "E20")

# ---------- Analysis: Expiring Soon ----------
ws2 = wb.create_sheet("Expiring_90d")
headers = ['customer','item','type','expire','days_left','price_year','contact','tel']
rows = []
for s, days in expiring_soon:
    rows.append([s.get('customer',''), s.get('item',''), s.get('type',''), s.get('expire',''), days, s.get('price_year',''), s.get('contact',''), s.get('tel','')])
write_table(ws2, headers, rows)

wb.save('/Users/nickamoto/ccl-crm-site/analysis/CCL_CRM_Analysis.xlsx')
print("SAVED")
print("total_revenue", total_revenue)
print("total_lost_revenue", total_lost_revenue)
print("expiring_soon", len(expiring_soon))
